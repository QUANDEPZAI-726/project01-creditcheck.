"""Offline: data_raw.xlsx -> data_encrypted.xlsx. Pseudonymization, not file encryption."""
import hashlib
import hmac
import json
import os
import re
import unicodedata
from pathlib import Path
from openpyxl import Workbook, load_workbook

FINANCIAL = [
    'Doanh thu thuần', 'Phải thu khách hàng', 'Giá vốn hàng bán',
    'Tài sản ngắn hạn', 'Tài sản cố định hữu hình thuần', 'Tổng tài sản',
    'Chi phí khấu hao và phân bổ', 'Chi phí bán hàng và quản lý doanh nghiệp',
    'Nợ ngắn hạn', 'Nợ dài hạn',
    'Lợi nhuận sau thuế từ hoạt động liên tục',
    'Lưu chuyển tiền thuần từ hoạt động kinh doanh',
]
RULES = {
    'Năm': 'KEEP',
    'Mã hồ sơ': 'HASH',
    'Tên người phụ trách hồ sơ': 'STAFF',
    'Tên doanh nghiệp': 'COMPANY',
    'Mã chứng khoán': 'COMPANY',
    'Ngành': 'DROP', 'Địa chỉ': 'DROP', 'Người đại diện': 'DROP',
    **{name: 'KEEP' for name in FINANCIAL},
}

def write_private(path, data):
    fd = os.open(path, os.O_WRONLY | os.O_CREAT | os.O_EXCL, 0o600)
    with os.fdopen(fd, 'wb') as f:
        f.write(data)

def protect(root=Path('.')):
    root = Path(root)
    src = load_workbook(root / 'data_raw.xlsx', read_only=True, data_only=False)
    if len(src.worksheets) != 1:
        raise ValueError('Script này chỉ hỗ trợ một sheet dữ liệu; cần chốt rule cho các sheet khác.')
    sheet = src.active
    data = list(sheet.iter_rows(values_only=True))
    headers, rows = list(data[0]), data[1:]
    if len(headers) != len(set(headers)) or set(headers) != set(RULES):
        raise ValueError('Các cột không khớp bảng quy tắc. Không tự giữ hoặc bỏ cột chưa được chốt.')
    pos = {name: i for i, name in enumerate(headers)}
    company_names, company_codes, pairs = {}, {}, set()
    for row in rows:
        if any(isinstance(v, str) and v.startswith('=') for v in row):
            raise ValueError('Có công thức Excel. Cần kiểm tra trước khi chuyển sang bản chia sẻ.')
        for name in ('Mã hồ sơ', 'Tên người phụ trách hồ sơ', 'Tên doanh nghiệp', 'Mã chứng khoán'):
            if not isinstance(row[pos[name]], str) or not row[pos[name]].strip():
                raise ValueError('Thiếu hoặc sai kiểu dữ liệu định danh; cần đối chiếu nguồn.')
        if not isinstance(row[pos['Năm']], int) or isinstance(row[pos['Năm']], bool):
            raise ValueError('Năm phải là số nguyên để liên kết các kỳ.')
        code = unicodedata.normalize('NFC', row[pos['Mã chứng khoán']]).strip().upper()
        name = unicodedata.normalize('NFC', row[pos['Tên doanh nghiệp']]).strip()
        if code in company_names and company_names[code] != name:
            raise ValueError('Một mã chứng khoán có nhiều tên. Chốt danh tính doanh nghiệp trước.')
        if name in company_codes and company_codes[name] != code:
            raise ValueError('Một doanh nghiệp có nhiều mã. Chốt mã ổn định trước khi xử lý.')
        company_names[code], company_codes[name] = name, code
        pair = (code, row[pos['Năm']])
        if pair in pairs:
            raise ValueError('Trùng doanh nghiệp–năm; cần đối chiếu, không tự gộp dữ liệu.')
        pairs.add(pair)

    private = root / '.privacy_private'
    private.mkdir(mode=0o700, exist_ok=True)
    key_path, map_path = private / 'hmac.key', private / 'staff_tokens.json'
    if not key_path.exists():
        if map_path.exists():
            raise ValueError('Đã có bảng giả danh nhưng thiếu khóa. Không tạo khóa mới làm đứt liên kết.')
        write_private(key_path, os.urandom(32))
    key = key_path.read_bytes()
    if len(key) != 32:
        raise ValueError('Khóa HMAC không hợp lệ.')
    staff_map = json.loads(map_path.read_text()) if map_path.exists() else {}
    if not isinstance(staff_map, dict) or len(set(staff_map.values())) != len(staff_map):
        raise ValueError('Bảng giả danh nhân viên không hợp lệ.')
    if any(not re.fullmatch(r'[0-9a-f]{64}', k) or not isinstance(v,str)
           or not re.fullmatch(r'STAFF_[0-9]{3,}',v) for k,v in staff_map.items()):
        raise ValueError('Bảng giả danh nhân viên sai định dạng.')

    def digest(domain, value):
        # HMAC có khóa, tách miền. Không dùng SHA-256 trực tiếp trên tên/mã dễ đoán.
        payload = json.dumps([domain, value], ensure_ascii=False).encode('utf-8')
        return hmac.new(key, payload, hashlib.sha256).hexdigest()

    tokens = {code: 'COMP_' + digest('company', code) for code in company_names}
    if len(set(tokens.values())) != len(tokens):
        raise ValueError('Phát hiện trùng token doanh nghiệp.')
    out = Workbook()
    ws = out.active
    ws.title = sheet.title
    keep_headers = [name for name in headers if RULES[name] != 'DROP']
    ws.append(keep_headers)
    for row in rows:
        code = unicodedata.normalize('NFC', row[pos['Mã chứng khoán']]).strip().upper()
        safe = []
        for name in keep_headers:
            value, rule = row[pos[name]], RULES[name]
            if rule == 'HASH':
                value = 'CASE_' + digest('case', value)
            elif rule == 'COMPANY':
                value = tokens[code]  # Tên và mã cùng dùng một token, không đưa năm vào hash.
            elif rule == 'STAFF':
                identity = unicodedata.normalize('NFC', value).strip()
                token = digest('staff', identity)
                if token not in staff_map:
                    used = set(staff_map.values())
                    n = 1
                    while f'STAFF_{n:03d}' in used:
                        n += 1
                    staff_map[token] = f'STAFF_{n:03d}'
                value = staff_map[token]
            safe.append(value)  # KEEP: giữ nguyên giá trị, không làm tròn hay đổi đơn vị.
        ws.append(safe)

    for old_row, new_row in zip(rows, list(ws.values)[1:]):
        for name in ['Năm'] + FINANCIAL:
            if old_row[pos[name]] != new_row[keep_headers.index(name)]:
                raise ValueError('Kiểm tra thất bại: năm hoặc số liệu tài chính đã thay đổi.')
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = ws.dimensions
    # Bảng riêng chỉ chứa HMAC -> STAFF, không lưu tên nhân viên dạng rõ.
    tmp_map = private / 'staff_tokens.tmp'
    fd = os.open(tmp_map, os.O_WRONLY | os.O_CREAT | os.O_TRUNC, 0o600)
    with os.fdopen(fd, 'w', encoding='utf-8') as f:
        json.dump(staff_map, f, ensure_ascii=False, indent=2)
    os.replace(tmp_map, map_path)
    target = root / 'data_encrypted.xlsx'
    out.save(target)
    src.close()
    print(f'Đã lưu {target.name}: {len(rows)} dòng, {len(keep_headers)} cột.')
    print('Năm và 12 trường tài chính giữ nguyên. Không chia sẻ thư mục .privacy_private.')

if __name__ == '__main__':
    protect()
